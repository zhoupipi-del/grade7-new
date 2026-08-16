"""
modules/research_ai/services_comment.py
=======================================
AI 学生评语生成核心业务逻辑。

隐私边界（延续作文模块设计）：
  - 学生真实姓名只存库、只在 docx 回填时使用；
  - 发送给 LLM 的 prompt 一律用「学生1/学生2」匿名编号，
    评语正文用「该同学」占位；
  - LLM 返回后，按 student_ref 映射回真实姓名并回填。

流程：取任务 → 构建匿名 prompt → 调 AI → 结构校验 → 业务校验 → 回填 → docx → 回写。
- Celery 异步（V2.2 新增）
"""
from __future__ import annotations

import json
import logging
import os
import time

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from sqlalchemy.ext.asyncio import AsyncSession

from core.models import get_local_now
from modules.research_ai.llm_client import (
    LLMError,
    LLMTransportError,
    chat_json,
    sanitize_student_info,
)
from modules.research_ai.models_comment import CommentTask
from modules.research_ai.schemas import CommentLLMResult
from modules.research_ai.word_renderer import (
    HEI,
    SONG,
    add_body,
    add_heading,
    add_label_body,
    add_paragraph,
    new_document,
    safe_filename_part,
)

logger = logging.getLogger(__name__)

OUTPUT_DIR = os.environ.get(
    "RESEARCH_AI_OUTPUT_DIR", "/opt/wings3/shared/uploads/research_ai"
)
COMMENTS_SUBDIR = "comments"

_TYPE_DESC = {
    "semester": "学期评语（全面回顾本学期的学习、行为、品德与成长，鼓励为主、具体客观）",
    "moral": "品德评语（侧重思想品德、行为习惯、集体意识、遵规守纪）",
    "parent": "家校沟通（面向家长，描述学生在校表现与进步，给出家校配合建议，语气温和积极）",
}


def _mask_known_names(text: str, names: set[str]) -> str:
    """把已知学生姓名定向替换为 [学生]（V2.2.1 P0 隐私闭合）。

    只替换长度≥2 的姓名（单字名如「王」易误伤「班主任」等常见词）。
    长名优先替换，避免短名被长名覆盖后残留。
    """
    if not text or not names:
        return text
    for name in sorted(names, key=len, reverse=True):
        if len(name) >= 2:
            text = text.replace(name, "[学生]")
    return text


def build_comment_prompt(task: CommentTask) -> str:
    """构建匿名评语 prompt。

    隐私边界（V2.2.1 P0 闭合）：学生真实姓名只存库、只用于 docx 回填；
    发送给 LLM 的 prompt 一律用「学生N」编号。keywords / style / class_name /
    term 等所有自由文本在拼入 prompt 前统一清洗：先定向替换本批次已知
    学生姓名为 [学生]，再过 sanitize_student_info() 清学号/手机号等数字 PII。
    """
    known_names = {str(s.get("name") or "").strip() for s in task.students}
    known_names = {n for n in known_names if n}

    def _clean(text: str) -> str:
        if not text:
            return ""
        return sanitize_student_info(_mask_known_names(str(text), known_names))

    lines = []
    for i, s in enumerate(task.students, 1):
        kw = _clean(s.get("keywords") or "")
        lines.append(f"学生{i}：{kw or '（无特别说明，按一般表现撰写）'}")
    student_list = "\n".join(lines)

    type_desc = _TYPE_DESC.get(task.comment_type, _TYPE_DESC["semester"])
    style = _clean(task.style) or "语言亲切自然，具体客观，多肯定、有期望"

    return f"""请为以下学生批量撰写{task.comment_type}。

【班级】{_clean(task.class_name)}
【学期】{_clean(task.term)}
【评语类型】{type_desc}
【风格要求】{style}

【学生名单（匿名编号）】
{student_list}

返回 JSON，必须严格使用以下结构：
{{
  "comments": [
    {{"student_ref": "学生1", "comment": "评语正文（用「该同学」指代学生本人，60-120字）"}},
    {{"student_ref": "学生2", "comment": "评语正文"}}
  ]
}}

规则：
1. 每个学生必须恰好一条评语，student_ref 与名单编号一一对应。
2. 评语正文用「该同学」指代，绝不出现真实姓名或编号。
3. 评语具体、真实、有针对性，结合给定关键词；不编造学生未提供的行为。
4. 语气鼓励为主，指出不足时用建设性措辞。
5. 只返回 JSON，不要任何额外文字。"""


def render_comment_docx(task: CommentTask, data: dict, out_path: str) -> None:
    """渲染批量评语 docx（回填真实姓名）。"""
    doc = new_document()

    title = f"{task.class_name}{task.term}学生评语"
    add_paragraph(
        doc, title, HEI, 22, True,
        align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=10,
    )

    # 姓名映射：studentN -> 真实姓名
    name_map = {f"学生{i}": s.get("name", "") for i, s in enumerate(task.students, 1)}
    # V2.2.1 P0：不信任 LLM 返回顺序——按 学生1..学生N 固定顺序回填。
    # ref 缺失/重复由 validator 在写入前拦截；此处 double-guard 用 .get 兜底，
    # 保证文档中每个学生恰好出现一次、顺序与教师输入一致。
    comments_by_ref = {
        str(item.get("student_ref", "")).strip(): item.get("comment", "")
        for item in data.get("comments", [])
    }

    for i in range(1, task.student_count + 1):
        ref = f"学生{i}"
        name = name_map.get(ref, ref)
        comment = comments_by_ref.get(ref, "")
        # 回填：评语中的「该同学」替换为真实姓名（教师本地展示用）
        comment = comment.replace("该同学", name)
        add_paragraph(doc, f"{i}. {name}", HEI, 13, True, before=6)
        add_body(doc, comment, indent=0.7)

    doc.save(out_path)


class CommentValidationError(ValueError):
    """评语业务校验失败。"""


async def generate_comments(db: AsyncSession, task_id: int) -> None:
    """生成学生评语并回写结果（Celery 异步任务调用）。"""
    task = await db.get(CommentTask, task_id)
    if not task:
        return
    task.status = "processing"
    db.add(task)
    await db.commit()

    t0 = time.time()
    try:
        prompt = build_comment_prompt(task)
        data = None

        for attempt in range(1, 3):
            data, usage, _elapsed = await chat_json(
                "你是一位经验丰富的中学班主任，评语写得具体、温暖、有针对性。只返回 JSON。",
                prompt,
                temperature=0.7,
            )
            try:
                result = CommentLLMResult(**data)
            except Exception as ve:
                logger.warning("[Comment] task=%s 结构校验失败(第%s次): %s", task_id, attempt, ve)
                if attempt < 2:
                    prompt += f"\n\n【重要纠错】你上次返回的JSON结构有误：{ve}，请修正后重新返回。"
                    continue
                raise CommentValidationError(f"JSON结构校验失败：{ve}") from ve

            biz_errors = result.validate_business_rules(task.student_count)
            if biz_errors:
                logger.warning("[Comment] task=%s 业务校验失败(第%s次): %s", task_id, attempt, biz_errors)
                if attempt < 2:
                    prompt += (
                        f"\n\n【重要纠错】你上次的评语有以下问题，请修正："
                        f"{'；'.join(biz_errors)}。必须每位学生一条且非空。"
                    )
                    continue
                raise CommentValidationError(f"业务规则校验失败：{'；'.join(biz_errors)}")

            break  # 校验通过

        out_dir = os.path.join(OUTPUT_DIR, COMMENTS_SUBDIR, str(task_id))
        os.makedirs(out_dir, exist_ok=True)
        # V2.2.1 P0：进入文件名的动态字段统一过 safe_filename_part
        # （term 如 "2025/2026学年" 含斜杠，不清洗会 FileNotFoundError）
        fname = (
            f"{safe_filename_part(task.class_name, '班级')}_"
            f"{safe_filename_part(task.term, '学期')}评语.docx"
        )
        out_path = os.path.join(out_dir, fname)
        render_comment_docx(task, data, out_path)

        rel_path = f"{COMMENTS_SUBDIR}/{task_id}/{fname}"
        task.status = "success"
        task.file_path = rel_path
        task.file_name = fname
        task.ai_raw_json = data
        task.prompt_tokens = usage.get("prompt_tokens")
        task.completion_tokens = usage.get("completion_tokens")
        task.elapsed_sec = int(time.time() - t0)
        task.finished_at = get_local_now()
        logger.info("[Comment] task=%s 生成成功，%s 人，耗时 %ss", task_id, task.student_count, task.elapsed_sec)

    except CommentValidationError as e:
        logger.error("[Comment] task=%s 校验失败: %s", task_id, e)
        task.status = "failed"
        task.error_msg = str(e)[:500]
        task.elapsed_sec = int(time.time() - t0)
        task.finished_at = get_local_now()
    except LLMError as e:
        logger.error("[Comment] task=%s LLM 错误: %s", task_id, e)
        task.status = "failed"
        task.error_msg = f"LLM 错误：{e}"[:500]
        task.elapsed_sec = int(time.time() - t0)
        task.finished_at = get_local_now()
    except LLMTransportError as e:
        # 传输错误：保持 processing，Celery 会重试，重试耗尽后 tasks.py 标记最终 failed
        logger.error("[Comment] task=%s 传输错误（Celery 将重试）: %s", task_id, e)
        task.status = "processing"
        task.error_msg = f"网络异常，正在自动重试：{e}"[:500]
        db.add(task)
        await db.commit()
        raise

    db.add(task)
    await db.commit()
