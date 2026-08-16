"""
modules/research_ai/services_analysis.py
=========================================
AI 学情分析报告核心业务逻辑。

流程：取任务 → 清洗数据摘要 → 调 AI → 结构校验 → 业务规则校验 → 生成 docx → 回写。
- 输入：用户提供的学情数据摘要（分数段分布/知识点正确率等）
- 隐私：摘要先过 sanitize_student_info()，防学号/手机号泄漏
- 输出 Word 学情报告（总体评估/数据解读/知识点掌握/学生分层/教学建议）
- Celery 异步（V2.2 新增）
"""
from __future__ import annotations

import logging
import os
import time

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from sqlalchemy.ext.asyncio import AsyncSession

from core.models import get_local_now
from modules.research_ai.llm_client import (
    LLMError,
    LLMTransportError,
    chat_json,
    sanitize_student_info,
)
from modules.research_ai.models_analysis import AnalysisTask
from modules.research_ai.schemas import AnalysisLLMResult
from modules.research_ai.word_renderer import (
    HEI,
    SONG,
    add_body,
    add_heading,
    add_label_body,
    add_paragraph,
    new_document,
    safe_filename_part,
)

logger = logging.getLogger(__name__)

OUTPUT_DIR = os.environ.get(
    "RESEARCH_AI_OUTPUT_DIR", "/opt/wings3/shared/uploads/research_ai"
)
ANALYSIS_SUBDIR = "analysis"


def build_analysis_prompt(task: AnalysisTask) -> str:
    """根据任务构建学情分析 prompt。

    隐私（V2.2.1 P0 闭合）：scope_name / exam_name / analysis_focus /
    data_summary 等所有进入 prompt 的自由文本字段统一过
    sanitize_student_info()（清学号/手机号等 8-12 位数字 PII），
    不再只清洗 data_summary 一个字段。
    """
    scope = sanitize_student_info(task.scope_name)
    exam = sanitize_student_info(task.exam_name or "")
    focus = sanitize_student_info(task.analysis_focus or "")
    summary = sanitize_student_info(task.data_summary)

    return f"""请基于以下学情数据，撰写一份专业的学情分析报告。

【分析对象】{scope}
【学科】{task.subject}
【年级】{task.grade}
【考试/阶段】{exam or '阶段测评'}
【学情数据摘要】
{summary}
【关注点】{focus or '整体学情与改进建议'}

返回 JSON，必须严格使用以下结构：
{{
  "overall_assessment": "总体评估（200字内，概括整体水平、优点与主要问题）",
  "score_distribution": {{"优秀率": "xx%", "及格率": "xx%", "低分率": "xx%", "解读": "分布特征解读"}},
  "findings": ["关键发现1", "关键发现2", "关键发现3"],
  "knowledge_mastery": [{{"knowledge_point": "知识点", "mastery": "掌握程度描述", "suggestion": "教学建议"}}],
  "student_stratification": [{{"layer": "学优生/中等生/学困生", "feature": "特征", "strategy": "提升策略"}}],
  "recommendations": ["教学建议1", "教学建议2", "教学建议3", "教学建议4", "教学建议5"]
}}

规则：
1. 所有结论必须基于给定数据，不得编造不存在的数据。
2. 不输出任何学生姓名、学号等个人信息（数据已脱敏）。
3. 教学建议至少 3 条，具体可操作。
4. 只返回 JSON，不要任何额外文字。"""


def render_analysis_docx(task: AnalysisTask, data: dict, out_path: str) -> None:
    """渲染学情分析报告 docx。"""
    doc = new_document()

    title = f"{task.scope_name}{task.subject}学情分析报告"
    add_paragraph(
        doc, title, HEI, 22, True,
        align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=4,
    )
    add_paragraph(
        doc, f"（{task.grade} · {task.exam_name or '阶段测评'}）",
        SONG, 12, False, align=WD_ALIGN_PARAGRAPH.CENTER, after=10,
    )

    add_heading(doc, "一、总体评估", level=1)
    add_body(doc, data.get("overall_assessment", ""))

    sd = data.get("score_distribution") or {}
    if sd:
        add_heading(doc, "二、分数段解读", level=1)
        for k, v in sd.items():
            if k == "解读":
                add_body(doc, f"解读：{v}")
            else:
                add_label_body(doc, f"{k}：", str(v))

    findings = data.get("findings") or []
    if findings:
        add_heading(doc, "三、关键发现", level=1)
        for i, f in enumerate(findings, 1):
            add_paragraph(doc, f"{i}. {f}", SONG, 12, indent=0.7)

    mastery = data.get("knowledge_mastery") or []
    if mastery:
        add_heading(doc, "四、知识点掌握情况", level=1)
        for i, m in enumerate(mastery, 1):
            add_label_body(
                doc,
                f"{i}. {m.get('knowledge_point', '')}：",
                f"{m.get('mastery', '')}（建议：{m.get('suggestion', '')}）",
            )

    strat = data.get("student_stratification") or []
    if strat:
        add_heading(doc, "五、学生分层建议", level=1)
        for i, s in enumerate(strat, 1):
            add_label_body(
                doc,
                f"{i}. {s.get('layer', '')}：",
                f"{s.get('feature', '')}；策略：{s.get('strategy', '')}",
            )

    recs = data.get("recommendations") or []
    if recs:
        add_heading(doc, "六、教学建议", level=1)
        for i, r in enumerate(recs, 1):
            add_paragraph(doc, f"{i}. {r}", SONG, 12, indent=0.7)

    add_paragraph(
        doc,
        "说明：本报告由 WINGS AI 教研助手基于所提供数据生成，仅供教师参考；"
        "请结合实际情况审阅后使用。",
        SONG, 10, before=12, indent=0.7,
    )

    doc.save(out_path)


class AnalysisValidationError(ValueError):
    """学情分析业务校验失败。"""


async def generate_analysis(db: AsyncSession, task_id: int) -> None:
    """生成学情分析报告并回写结果（Celery 异步任务调用）。"""
    task = await db.get(AnalysisTask, task_id)
    if not task:
        return
    task.status = "processing"
    db.add(task)
    await db.commit()

    t0 = time.time()
    try:
        # V2.2.1 P0：隐私清洗已内聚到 build_analysis_prompt()
        # （scope_name/exam_name/analysis_focus/data_summary 全部过 sanitizer）
        prompt = build_analysis_prompt(task)
        data = None

        for attempt in range(1, 3):
            data, usage, _elapsed = await chat_json(
                "你是一位资深教务分析专家，擅长从成绩数据中提炼学情特征，"
                "给出客观、专业、可操作的教学建议。只返回 JSON。",
                prompt,
                temperature=0.5,
            )
            try:
                result = AnalysisLLMResult(**data)
            except Exception as ve:
                logger.warning("[Analysis] task=%s 结构校验失败(第%s次): %s", task_id, attempt, ve)
                if attempt < 2:
                    prompt += f"\n\n【重要纠错】你上次返回的JSON结构有误：{ve}，请修正后重新返回。"
                    continue
                raise AnalysisValidationError(f"JSON结构校验失败：{ve}") from ve

            biz_errors = result.validate_business_rules()
            if biz_errors:
                logger.warning("[Analysis] task=%s 业务校验失败(第%s次): %s", task_id, attempt, biz_errors)
                if attempt < 2:
                    prompt += f"\n\n【重要纠错】你上次的报告有以下问题，请修正：{'；'.join(biz_errors)}"
                    continue
                raise AnalysisValidationError(f"业务规则校验失败：{'；'.join(biz_errors)}")

            break  # 校验通过

        out_dir = os.path.join(OUTPUT_DIR, ANALYSIS_SUBDIR, str(task_id))
        os.makedirs(out_dir, exist_ok=True)
        # V2.2.1 P0：进入文件名的动态字段统一过 safe_filename_part
        fname = (
            f"{safe_filename_part(task.scope_name, '学情')}_"
            f"{safe_filename_part(task.subject, '学科')}学情分析.docx"
        )
        out_path = os.path.join(out_dir, fname)
        render_analysis_docx(task, data, out_path)

        rel_path = f"{ANALYSIS_SUBDIR}/{task_id}/{fname}"
        task.status = "success"
        task.file_path = rel_path
        task.file_name = fname
        task.ai_raw_json = data
        task.prompt_tokens = usage.get("prompt_tokens")
        task.completion_tokens = usage.get("completion_tokens")
        task.elapsed_sec = int(time.time() - t0)
        task.finished_at = get_local_now()
        logger.info("[Analysis] task=%s 生成成功，耗时 %ss", task_id, task.elapsed_sec)

    except AnalysisValidationError as e:
        logger.error("[Analysis] task=%s 校验失败: %s", task_id, e)
        task.status = "failed"
        task.error_msg = str(e)[:500]
        task.elapsed_sec = int(time.time() - t0)
        task.finished_at = get_local_now()
    except LLMError as e:
        logger.error("[Analysis] task=%s LLM 错误: %s", task_id, e)
        task.status = "failed"
        task.error_msg = f"LLM 错误：{e}"[:500]
        task.elapsed_sec = int(time.time() - t0)
        task.finished_at = get_local_now()
    except LLMTransportError as e:
        # 传输错误：保持 processing，Celery 会重试，重试耗尽后 tasks.py 标记最终 failed
        logger.error("[Analysis] task=%s 传输错误（Celery 将重试）: %s", task_id, e)
        task.status = "processing"
        task.error_msg = f"网络异常，正在自动重试：{e}"[:500]
        db.add(task)
        await db.commit()
        raise

    db.add(task)
    await db.commit()
