"""
modules/research_ai/services_lesson_plan.py
============================================
批量教案生成核心业务逻辑。

V2.1 修复：
  - [P1] LLM 结果业务校验（教案总时长=periods×45，数学素养3~4个），校验失败重试一次
  - [P1] 批次状态聚合竞态修复：SELECT ... FOR UPDATE 串行化
  - [P1] 异常不再被静默吞掉后 return，区分 LLM 错误（标记 failed）和其他错误（可 Celery 重试）
  - 输出目录修正为 /opt/wings3/shared/uploads/research_ai/
"""
from __future__ import annotations

import logging
import os
import re
import time

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from core.models import get_local_now
from modules.research_ai.llm_client import LLMError, LLMTransportError, chat_json
from modules.research_ai.models_lesson_plan import LessonPlanItem, LessonPlanTask
from modules.research_ai.schemas import LessonLLMResult
from modules.research_ai.word_renderer import (
    HEI,
    SONG,
    add_heading,
    add_label_body,
    add_paragraph,
    fill_cell,
    make_table,
    new_document,
    set_font,
    shade_cell,
)

logger = logging.getLogger(__name__)

OUTPUT_DIR = os.environ.get(
    "RESEARCH_AI_OUTPUT_DIR", "/opt/wings3/shared/uploads/research_ai"
)
LESSON_PLANS_SUBDIR = "lesson_plans"

# ════════════════════════════════════════════════════════════
# 学科核心素养框架
# ════════════════════════════════════════════════════════════
SUBJECT_FRAMEWORKS: dict[str, list[str]] = {
    "语文": ["文化自信", "语言运用", "思维能力", "审美创造"],
    "数学": [
        "数感", "量感", "符号意识", "运算能力", "几何直观",
        "空间观念", "推理意识", "数据意识", "模型意识", "应用意识", "创新意识",
    ],
    "英语": ["语言能力", "文化意识", "思维品质", "学习能力"],
    "物理": ["物理观念", "科学思维", "探究实践", "态度责任"],
    "化学": ["化学观念", "科学思维", "探究实践", "态度责任"],
    "生物": ["生命观念", "科学思维", "探究实践", "态度责任"],
    "道德与法治": ["政治认同", "道德修养", "法治观念", "健全人格", "责任意识"],
    "道法": ["政治认同", "道德修养", "法治观念", "健全人格", "责任意识"],
    "历史": ["唯物史观", "时空观念", "史料实证", "历史解释", "家国情怀"],
    "地理": ["人地协调观", "综合思维", "区域认知", "地理实践力"],
}
GENERIC_FRAMEWORK = [
    "知识与技能", "过程与方法", "情感态度与价值观", "核心素养发展"
]


def get_framework(subject: str) -> list[str]:
    s = (subject or "").strip()
    for key, fw in SUBJECT_FRAMEWORKS.items():
        if key in s:
            return fw
    return GENERIC_FRAMEWORK


def curriculum_standard_hint(subject: str) -> str:
    s = (subject or "").strip()
    hints = {
        "语文": "依据《义务教育语文课程标准（2022年版）》相应学段学业要求与核心素养。",
        "数学": "依据《义务教育数学课程标准（2022年版）》相应学段课程内容与学业要求。",
        "英语": "依据《义务教育英语课程标准（2022年版）》相应级别课程目标与核心素养。",
        "物理": "依据《义务教育物理课程标准（2022年版）》相应主题内容与学业要求。",
        "化学": "依据《义务教育化学课程标准（2022年版）》相应主题内容与学业要求。",
        "生物": "依据《义务教育生物学课程标准（2022年版）》相应学习主题与学业要求。",
        "历史": "依据《义务教育历史课程标准（2022年版）》相应板块内容与学业要求。",
        "地理": "依据《义务教育地理课程标准（2022年版）》相应主题内容与学业要求。",
    }
    for k, v in hints.items():
        if k in s:
            return v
    if "道法" in s or "道德与法治" in s:
        return "依据《义务教育道德与法治课程标准（2022年版）》相应学段内容要求与核心素养。"
    return "依据《义务教育课程方案和课程标准（2022年版）》相应学科课程目标与学业要求。"


SYSTEM_PROMPT = (
    "你是一位经验丰富的初中各学科教研组长与特级教师，精通《义务教育课程方案和课程标准"
    "（2022年版）》，擅长撰写规范、详实、可直接用于课堂的教学设计。"
    "你的输出必须是严格的 JSON（不要 markdown、不要解释性文字）。"
)


def build_lesson_prompt(item: LessonPlanItem) -> str:
    framework = get_framework(item.subject)
    is_math = "数学" in item.subject
    obj_instr = (
        "数学学科请从该列表中按课题紧密选取 3-4 个最相关的核心素养作为目标，"
        "不要全部罗列，每个目标要写具体可观测的行为表现。"
        if is_math
        else "请围绕该学科的核心素养框架逐一撰写教学目标，每个维度要具体、可操作、贴合本课题。"
    )
    kp = (
        f"教师指定教学重点：{item.key_point}"
        if item.key_point
        else "教学重点：由你根据课题推断"
    )
    dp = (
        f"教师指定教学难点：{item.diff_point}"
        if item.diff_point
        else "教学难点：由你根据课题推断"
    )

    return f"""请为以下一节课生成一份完整规范的教学设计，严格返回 JSON。

【课程基本信息】
学科：{item.subject}
教材版本：{item.textbook or ""}
年级：{item.grade}
单元/章节：{item.unit or ""}
课题：{item.topic}
课时：{item.periods} 课时（每课时 45 分钟）
{kp}
{dp}

{curriculum_standard_hint(item.subject)}
该学科核心素养框架：{framework}
{obj_instr}

【JSON 结构要求】（字段名必须完全一致，所有值用中文字符串；环节/步骤要具体可执行）
{{
  "课标依据": "字符串，引用 2022 版课标相应主题的具体要求，2-4 条",
  "教学目标": {{"维度一名称": "对应目标的具体描述", "维度二名称": "..."}},
  "教学重点": "字符串，分点列出",
  "教学难点": "字符串，分点列出",
  "教学方法": "字符串，列出主要教学方法并简述如何组合使用",
  "教学准备": {{"教师准备": "字符串", "学生准备": "字符串"}},
  "教学过程": [
    {{"环节": "如 导入新课", "时长": "如 5分钟",
      "教师活动": ["步骤1..."], "学生活动": ["步骤1..."], "设计意图": "字符串"}}
  ],
  "板书设计": "结构化板书，用换行和缩进，不要写'略'",
  "作业布置": {{"基础": "字符串", "提升": "字符串", "拓展": "字符串"}},
  "教学反思提示": ["提示问题1", "提示问题2", "提示问题3", "提示问题4"]
}}

【教学过程要求】
1. 按"导入→新授→练习→小结→作业"逻辑组织；
2. 所有环节时长之和必须等于 {item.periods}×45 分钟（共{item.periods * 45}分钟）；
3. 教师/学生活动要分条、具体可执行；
4. 每个环节写清设计意图。

只返回 JSON，不要任何额外文字。"""


# ════════════════════════════════════════════════════════════
# Word 渲染
# ════════════════════════════════════════════════════════════

def _sanitize(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", str(name))
    name = re.sub(r"\s+", "", name)
    return name or "未命名"


def render_lesson_docx(item: LessonPlanItem, data: dict, out_path: str) -> None:
    doc = new_document()

    add_paragraph(
        doc, "教  学  设  计", HEI, 22, True,
        align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6,
    )
    add_paragraph(
        doc, f"——《{item.topic}》", HEI, 16, True,
        align=WD_ALIGN_PARAGRAPH.CENTER, after=12,
    )

    add_heading(doc, "一、基本信息", level=1)
    info = make_table(doc, rows=4, cols=4)
    rows = [
        ["学    科", item.subject, "年    级", item.grade],
        ["教    材", item.textbook or "", "课    时", f"{item.periods}课时（45分钟/课时）"],
        ["课    题", item.topic, "单元位置", item.unit or ""],
    ]
    for i, rd in enumerate(rows):
        for j, txt in enumerate(rd):
            cell = info.rows[i].cells[j]
            if j in (0, 2):
                fill_cell(cell, txt, HEI, 11, True,
                          WD_ALIGN_PARAGRAPH.CENTER, shade="F0F0F0")
            else:
                fill_cell(cell, txt, SONG, 11)
    merged = info.rows[3].cells[0]
    for k in (1, 2, 3):
        merged = merged.merge(info.rows[3].cells[k])
    fill_cell(
        merged,
        "说明：本教案由 WINGS AI 教研助手辅助生成，供教师参考；请结合学情审阅修改后使用。",
        SONG, 10, shade="FAFAFA",
    )

    add_heading(doc, "二、课标依据", level=1)
    cb = data.get("课标依据", "")
    if isinstance(cb, list):
        for x in cb:
            add_paragraph(doc, f"· {x}", SONG, 11, indent=0.7)
    else:
        for line in str(cb).split("\n"):
            line = line.strip()
            if line:
                add_paragraph(doc, line, SONG, 11, indent=0.7)

    add_heading(doc, "三、教学目标（核心素养）", level=1)
    objs = data.get("教学目标", {})
    if isinstance(objs, dict):
        for k, v in objs.items():
            add_label_body(doc, f"{k}：", str(v))
    elif isinstance(objs, list):
        for x in objs:
            add_paragraph(doc, f"· {x}", SONG, 11, indent=0.7)

    add_heading(doc, "四、教学重难点", level=1)
    add_label_body(doc, "教学重点：", str(data.get("教学重点", "")))
    add_label_body(doc, "教学难点：", str(data.get("教学难点", "")))
    if item.key_point:
        add_label_body(doc, "教师指定重点：", item.key_point)
    if item.diff_point:
        add_label_body(doc, "教师指定难点：", item.diff_point)

    add_heading(doc, "五、教学方法", level=1)
    add_paragraph(doc, str(data.get("教学方法", "")), SONG, 11, indent=0.7)

    add_heading(doc, "六、教学准备", level=1)
    prep = data.get("教学准备", {})
    if isinstance(prep, dict):
        add_label_body(doc, "教师准备：", str(prep.get("教师准备", "")))
        add_label_body(doc, "学生准备：", str(prep.get("学生准备", "")))

    add_heading(doc, "七、教学过程", level=1)
    tp = make_table(doc, rows=1, cols=3)
    headers = ["教学环节", "教师活动 / 学生活动", "设计意图"]
    for i, h in enumerate(headers):
        fill_cell(tp.rows[0].cells[i], h, HEI, 11, True,
                  WD_ALIGN_PARAGRAPH.CENTER, shade="D5E8F0")
    widths = [Cm(2.8), Cm(8.5), Cm(4.7)]
    for i, w in enumerate(widths):
        tp.rows[0].cells[i].width = w

    for stage in data.get("教学过程", []) or []:
        row = tp.add_row()
        for i, w in enumerate(widths):
            row.cells[i].width = w

        c0 = row.cells[0]
        c0.text = ""
        title = stage.get("环节", "")
        dur = stage.get("时长", "")
        p = c0.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{title}\n（{dur}）" if dur else title)
        set_font(r, HEI, 10, True)

        c1 = row.cells[1]
        c1.text = ""
        p = c1.paragraphs[0]
        r = p.add_run("【教师活动】")
        set_font(r, HEI, 9.5, True)
        for act in stage.get("教师活动", []) or []:
            ap = c1.add_paragraph()
            ar = ap.add_run(str(act))
            set_font(ar, SONG, 9.5)
        sp = c1.add_paragraph()
        sr = sp.add_run("【学生活动】")
        set_font(sr, HEI, 9.5, True)
        for act in stage.get("学生活动", []) or []:
            ap = c1.add_paragraph()
            ar = ap.add_run(str(act))
            set_font(ar, SONG, 9.5)

        c2 = row.cells[2]
        c2.text = ""
        p = c2.paragraphs[0]
        r = p.add_run(str(stage.get("设计意图", "")))
        set_font(r, SONG, 9.5)

    add_heading(doc, "八、板书设计", level=1)
    bt = make_table(doc, rows=1, cols=1)
    bc = bt.rows[0].cells[0]
    bc.text = ""
    board = str(data.get("板书设计", "（略）"))
    for i, line in enumerate(board.split("\n")):
        p = bc.paragraphs[0] if i == 0 else bc.add_paragraph()
        r = p.add_run(line or " ")
        set_font(r, SONG, 11)

    add_heading(doc, "九、作业布置（分层）", level=1)
    hw = data.get("作业布置", {})
    if isinstance(hw, dict):
        add_label_body(doc, "基础层（全体）：", str(hw.get("基础", "")))
        add_label_body(doc, "提升层（中等）：", str(hw.get("提升", "")))
        add_label_body(doc, "拓展层（学有余力）：", str(hw.get("拓展", "")))

    add_heading(doc, "十、教学反思（教师课后填写）", level=1)
    tips = data.get("教学反思提示", []) or []
    if tips:
        for i, t in enumerate(tips, 1):
            add_paragraph(doc, f"{i}. {t}", HEI, 11, True, before=4)
            add_paragraph(doc, "（课后填写）", SONG, 11, indent=0.7, after=4)
    else:
        for i in range(1, 5):
            add_paragraph(doc, f"{i}. 反思要点：", HEI, 11, True, before=4)
            add_paragraph(doc, "（课后填写）", SONG, 11, indent=0.7, after=4)

    add_paragraph(
        doc, "本教案由 AI 辅助生成，最终内容以教师审阅修改为准。",
        SONG, 9, align=WD_ALIGN_PARAGRAPH.CENTER, before=12,
    )
    doc.save(out_path)


# ════════════════════════════════════════════════════════════
# 单课生成
# ════════════════════════════════════════════════════════════

class LessonValidationError(ValueError):
    """教案 LLM 结果业务校验失败。"""


async def generate_one_lesson(db: AsyncSession, item_id: int) -> None:
    """取 item → 调 AI → 校验 → 生成 docx → 回写 DB。"""
    item = await db.get(LessonPlanItem, item_id)
    if not item:
        return
    item.status = "processing"
    db.add(item)
    await db.commit()

    t0 = time.time()
    try:
        prompt = build_lesson_prompt(item)
        data = None

        # [P1] 业务校验失败重试一次
        for attempt in range(1, 3):
            data, usage, _elapsed = await chat_json(
                SYSTEM_PROMPT, prompt, temperature=0.5
            )

            # 结构校验
            try:
                result = LessonLLMResult(**data)
            except Exception as ve:
                logger.warning(
                    "[LessonPlan] item=%s 结构校验失败(第%s次): %s",
                    item_id, attempt, ve,
                )
                if attempt < 2:
                    prompt += f"\n\n【重要纠错】你上次返回的JSON结构有误：{ve}，请修正后重新返回。"
                    continue
                raise LessonValidationError(f"JSON结构校验失败：{ve}") from ve

            # 业务规则校验
            biz_errors = result.validate_business_rules(item.periods, item.subject)
            if biz_errors:
                logger.warning(
                    "[LessonPlan] item=%s 业务校验失败(第%s次): %s",
                    item_id, attempt, biz_errors,
                )
                if attempt < 2:
                    prompt += (
                        f"\n\n【重要纠错】你上次的教案有以下问题，请修正："
                        f"{'；'.join(biz_errors)}。请确保各环节时长之和={item.periods * 45}分钟。"
                    )
                    continue
                raise LessonValidationError(f"业务规则校验失败：{'；'.join(biz_errors)}")

            break  # 校验通过

        # 教师指定的重难点覆盖 AI
        if item.key_point:
            data["教学重点"] = item.key_point
        if item.diff_point:
            data["教学难点"] = item.diff_point

        out_dir = os.path.join(OUTPUT_DIR, LESSON_PLANS_SUBDIR, str(item.task_id))
        os.makedirs(out_dir, exist_ok=True)
        fname = (
            f"{_sanitize(item.subject)}_{_sanitize(item.grade)}"
            f"_{_sanitize(item.topic)}.docx"
        )
        out_path = os.path.join(out_dir, fname)
        render_lesson_docx(item, data, out_path)

        rel_path = f"{LESSON_PLANS_SUBDIR}/{item.task_id}/{fname}"
        item.status = "success"
        item.file_path = rel_path
        item.file_name = fname
        item.ai_raw_json = data
        item.prompt_tokens = usage.get("prompt_tokens")
        item.completion_tokens = usage.get("completion_tokens")
        item.elapsed_sec = int(time.time() - t0)
        item.finished_at = get_local_now()
        logger.info(
            "[LessonPlan] item=%s 生成成功，耗时 %ss", item_id, item.elapsed_sec
        )
    except LessonValidationError as e:
        # 业务校验失败：标记 failed，不重试（Celery 不会重试，因为我们正常 return）
        logger.error("[LessonPlan] item=%s 校验失败: %s", item_id, e)
        item.status = "failed"
        item.error_msg = str(e)[:500]
        item.elapsed_sec = int(time.time() - t0)
        item.finished_at = get_local_now()
    except LLMError as e:
        # LLM 配置/JSON 解析错误：标记 failed
        logger.error("[LessonPlan] item=%s LLM 错误: %s", item_id, e)
        item.status = "failed"
        item.error_msg = f"LLM 错误：{e}"[:500]
        item.elapsed_sec = int(time.time() - t0)
        item.finished_at = get_local_now()
    except LLMTransportError as e:
        # [V2.1.1] 传输错误：保持 processing，不写 terminal failed，
        # 让前端继续轮询；Celery 会重试，重试耗尽后由 tasks.py 标记最终 failed。
        logger.error("[LessonPlan] item=%s 传输错误（Celery 将重试）: %s", item_id, e)
        item.status = "processing"
        item.error_msg = f"网络异常，正在自动重试：{e}"[:500]
        db.add(item)
        await db.commit()
        await _refresh_task_status(db, item.task_id)
        raise  # 向上抛出让 Celery retry

    db.add(item)
    await db.commit()
    await _refresh_task_status(db, item.task_id)


async def _refresh_task_status(db: AsyncSession, task_id: int) -> None:
    """
    聚合刷新任务状态与计数。
    [P1 修复] 使用 SELECT ... FOR UPDATE 锁定 task 行，
    防止多个 worker 并发聚合时状态写回竞态。
    """
    # 锁定 task 行（事务内行锁，提交后释放）
    result = await db.execute(
        select(LessonPlanTask)
        .where(LessonPlanTask.id == task_id)
        .with_for_update()
    )
    task = result.scalar_one_or_none()
    if not task:
        return

    # 在锁内重新读取最新 item 状态
    items_result = await db.execute(
        select(LessonPlanItem.status).where(LessonPlanItem.task_id == task_id)
    )
    statuses = [r[0] for r in items_result.all()]

    task.success_count = sum(1 for s in statuses if s == "success")
    task.fail_count = sum(1 for s in statuses if s == "failed")
    pending = sum(1 for s in statuses if s in ("pending", "processing"))

    if pending == 0:
        if task.fail_count == 0:
            task.status = "completed"
        elif task.success_count == 0:
            task.status = "failed"
        else:
            task.status = "partial_failed"
        task.finished_at = get_local_now()
    else:
        task.status = "processing"
    db.add(task)
    await db.commit()
